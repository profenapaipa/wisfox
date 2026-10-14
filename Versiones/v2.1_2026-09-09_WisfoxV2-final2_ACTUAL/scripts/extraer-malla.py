import openpyxl, json, re, sys
wb = openpyxl.load_workbook('MALLA INGLES 2026.xlsx', data_only=True)

GRADOS = {'Sixth':'6','Séptimo':'7','Octavo':'8','Noveno':'9','Décimo':'10','Once':'11'}
ROMAN = {'I':1,'II':2,'III':3,'IV':4}

def g(ws, coord):
    v = ws[coord].value
    if v is None: return ''
    return str(v).strip()

out = {}
for name in wb.sheetnames:
    parts = name.strip().split(' ',1)
    if len(parts)!=2: continue
    rom, grado_nom = parts[0], parts[1].strip()
    if rom not in ROMAN or grado_nom not in GRADOS: continue
    ws = wb[name]
    gr = GRADOS[grado_nom]
    per = ROMAN[rom]
    d = {
      'sheet': name,
      'grado': gr,
      'grado_nombre': g(ws,'C6'),
      'periodo': per,
      'intensidad_horaria': g(ws,'I6'),
      'semanas': g(ws,'L6'),
      'total_horas': g(ws,'O6'),
      'estandares': g(ws,'D7'),
      'competencias': g(ws,'E11'),
      'componentes': g(ws,'M11'),
      'concepto_tematico': g(ws,'E16'),
      'afirmacion_1': g(ws,'B21'),
      'afirmacion_2': g(ws,'B22'),
      'evidencia_1': g(ws,'G21'),
      'evidencia_2': g(ws,'G22'),
    }
    out.setdefault(gr, {})[str(per)] = d

json.dump(out, open(sys.argv[1],'w',encoding='utf-8'), ensure_ascii=False, indent=1)
print('grados:', sorted(out.keys(), key=int))
for gr in sorted(out, key=int):
    print(gr, sorted(out[gr].keys()), [out[gr][p]['concepto_tematico'][:60] for p in sorted(out[gr])])
